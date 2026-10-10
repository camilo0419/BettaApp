"""Valida e importa el catálogo UNSPSC oficial por lotes."""

import csv
import hashlib
import io
import re
import shutil
import tempfile
import time
import unicodedata
import zipfile
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from tienda.models import SyncAuditLog, UNSPSCCode


CONFIRMATION = "IMPORTAR CATÁLOGO UNSPSC"
MAX_FILE_BYTES = 400 * 1024 * 1024
MAX_ZIP_ENTRIES = 20
MAX_ZIP_COMPRESSION_RATIO = 1000
COPY_CHUNK_BYTES = 1024 * 1024
LEVEL_FIELDS = {
    "segment": ("Segmento", "Título del Segmento"),
    "family": ("Familia", "Título de Familia"),
    "class": ("Clase", "Título de Clase"),
    "product": ("Producto", "Título Producto"),
}


def _norm(value):
    value = unicodedata.normalize("NFKD", str(value or "").strip().casefold())
    value = "".join(char for char in value if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]", "", value)


def _code(value):
    text = str(value or "").strip()
    if text.endswith(".0"):
        text = text[:-2]
    digits = re.sub(r"\D", "", text)
    if len(digits) == 7:
        digits = digits.zfill(8)
    return digits if len(digits) == 8 else ""


def _is_active(value):
    return str(value or "").strip().casefold() not in {"0", "no", "false", "inactivo", "deshabilitado", "disabled"}


def _header_spec(headers):
    """Detecta una estructura UNSPSC sin exigir que el encabezado sea la fila 1."""
    normalized = {_norm(header): header for header in headers if str(header or "").strip()}
    hierarchy = {
        level: (normalized.get(_norm(code)), normalized.get(_norm(title)))
        for level, (code, title) in LEVEL_FIELDS.items()
    }
    if hierarchy["segment"][0] and hierarchy["segment"][1] and hierarchy["product"][0] and hierarchy["product"][1]:
        return {"kind": "hierarchy", "fields": hierarchy}
    code_header = next((header for key, header in normalized.items() if key in {"codigo", "code", "codigounspsc", "unspsc"}), None)
    description_header = next((header for key, header in normalized.items() if key in {"descripcion", "description", "nombre", "producto"}), None)
    active_header = next((header for key, header in normalized.items() if key in {"vigente", "activo", "active", "estado", "status"}), None)
    if code_header and description_header:
        return {"kind": "flat", "code": code_header, "description": description_header, "active": active_header}
    return None


def _record_from_row(row, spec, catalog_version, source, source_row):
    def value(header):
        return row.get(header) if header else None

    if spec["kind"] == "hierarchy":
        for level, (code_header, title_header) in spec["fields"].items():
            code = _code(value(code_header))
            title = str(value(title_header) or "").strip()
            if code and title:
                yield {"code": code, "description": title[:255], "level": level, "active": True, "catalog_version": catalog_version, "source": source, "source_row": source_row}
        return
    code = _code(value(spec["code"]))
    description = str(value(spec["description"]) or "").strip()
    if code and description:
        level = "segment" if code.endswith("000000") else "family" if code.endswith("0000") else "class" if code.endswith("00") else "product"
        yield {"code": code, "description": description[:255], "level": level, "active": _is_active(value(spec["active"])) if spec["active"] else True, "catalog_version": catalog_version, "source": source, "source_row": source_row}


def _source_file(path):
    path = Path(path).resolve()
    if not path.exists() or not path.is_file():
        raise CommandError("El archivo UNSPSC no existe o no es un archivo regular.")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise CommandError("El archivo UNSPSC supera el límite permitido de 400 MB.")
    if path.suffix.casefold() != ".zip":
        return path
    with zipfile.ZipFile(path) as archive:
        entries = [info for info in archive.infolist() if not info.is_dir()]
        if len(entries) > MAX_ZIP_ENTRIES:
            raise CommandError("El ZIP contiene demasiados archivos.")
        files = [info for info in entries if Path(info.filename).suffix.casefold() in {".csv", ".xlsx"}]
        if len(entries) != 1 or len(files) != 1:
            raise CommandError("El ZIP debe contener exactamente un CSV o XLSX de catálogo.")
        info = files[0]
        member_path = Path(info.filename)
        if member_path.is_absolute() or ".." in member_path.parts:
            raise CommandError("El ZIP contiene una ruta de archivo no segura.")
        if info.file_size > MAX_FILE_BYTES or info.compress_size == 0 or info.file_size / max(info.compress_size, 1) > MAX_ZIP_COMPRESSION_RATIO:
            raise CommandError("El ZIP no supera las validaciones de tamaño o compresión.")
        target = Path(tempfile.gettempdir()) / f"unspsc-source-{time.time_ns()}{member_path.suffix.casefold()}"
        with target.open("wb") as handle:
            with archive.open(info, "r") as source:
                shutil.copyfileobj(source, handle, length=COPY_CHUNK_BYTES)
        return target


def _file_sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(COPY_CHUNK_BYTES), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _iter_csv(path, catalog_version, source):
    raw_handle = path.open("rb")
    try:
        probe = raw_handle.read(65536)
        raw_handle.seek(0)
        try:
            probe.decode("utf-8-sig")
            encoding = "utf-8-sig"
        except UnicodeDecodeError:
            encoding = "cp1252"
        with io.TextIOWrapper(raw_handle, encoding=encoding, newline="") as handle:
            sample = handle.read(8192)
            handle.seek(0)
            try:
                delimiter = csv.Sniffer().sniff(sample, delimiters=";,\t").delimiter
            except csv.Error:
                delimiter = ";"
            reader = csv.reader(handle, delimiter=delimiter)
            spec = None
            headers = None
            for source_row, values in enumerate(reader, start=1):
                candidate = _header_spec(values)
                if candidate:
                    headers, spec = values, candidate
                    break
                if source_row >= 50:
                    break
            if not spec:
                raise CommandError("No se identificaron columnas UNSPSC válidas en las primeras 50 filas.")
            for source_row, values in enumerate(reader, start=source_row + 1):
                row = dict(zip(headers, values))
                yield from _record_from_row(row, spec, catalog_version, source, source_row)
    finally:
        if not raw_handle.closed:
            raw_handle.close()


def _iter_rows(path, catalog_version, source):
    source_path = _source_file(path)
    try:
        if source_path.suffix.casefold() == ".csv":
            yield from _iter_csv(source_path, catalog_version, source)
            return
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise CommandError("El runtime no tiene openpyxl para leer el XLSX oficial.") from exc
        workbook = load_workbook(source_path, read_only=True, data_only=True)
        try:
            found = False
            for sheet in workbook.worksheets:
                values = sheet.iter_rows(values_only=True)
                buffered = []
                spec = None
                headers = None
                for source_row, values_row in enumerate(values, start=1):
                    buffered.append(values_row)
                    candidate_headers = [str(value or "").strip() for value in values_row]
                    spec = _header_spec(candidate_headers)
                    if spec:
                        headers = candidate_headers
                        found = True
                        break
                    if source_row >= 50:
                        break
                if not spec:
                    continue
                for source_row, values_row in enumerate(values, start=source_row + 1):
                    row = dict(zip(headers, values_row))
                    yield from _record_from_row(row, spec, catalog_version, source, source_row)
                break
            if not found:
                raise CommandError("No se identificaron columnas UNSPSC válidas en las primeras 50 filas de las hojas.")
        finally:
            workbook.close()
    finally:
        if source_path != Path(path).resolve() and source_path.exists():
            source_path.unlink(missing_ok=True)


def _parent_code(code):
    if code.endswith("000000"):
        return None
    if code.endswith("0000"):
        return code[:2] + "000000"
    if code.endswith("00"):
        return code[:4] + "0000"
    return code[:6] + "00"


class Command(BaseCommand):
    help = "Valida/importa un catálogo UNSPSC oficial; dry-run por defecto."

    def add_arguments(self, parser):
        parser.add_argument("--file", required=True, help="Ruta local al CSV, XLSX o ZIP oficial.")
        parser.add_argument("--catalog-version", required=True, help="Versión efectiva declarada por la fuente.")
        parser.add_argument("--source-url", default="", help="URL oficial de origen.")
        parser.add_argument("--batch-size", type=int, default=1000)
        parser.add_argument("--apply", action="store_true")
        parser.add_argument("--confirm", default="")
        parser.add_argument("--expected-fingerprint", default="", help="Huella validada previamente; bloquea archivos modificados.")

    def handle(self, *args, **options):
        batch_size = max(1, min(options["batch_size"], 5000))
        source = options["source_url"]
        fingerprint = _file_sha256(options["file"])
        expected = str(options.get("expected_fingerprint") or "").strip()
        if expected and expected != fingerprint:
            raise CommandError("El archivo cambió después de la vista previa; se requiere una nueva validación.")
        seen = {}
        duplicates = invalid = rows = source_rows = 0
        last_source_row = None
        by_level = {key: 0 for key in LEVEL_FIELDS}
        started = time.monotonic()
        for record in _iter_rows(options["file"], options["catalog_version"], source):
            rows += 1
            if record.get("source_row") != last_source_row:
                source_rows += 1
                last_source_row = record.get("source_row")
            code = record["code"]
            if code in seen:
                duplicates += 1
                continue
            if len(code) != 8 or not code.isdigit():
                invalid += 1
                continue
            seen[code] = record
            by_level[record["level"]] += 1
        elapsed = round(time.monotonic() - started, 2)
        if not seen:
            raise CommandError("El archivo no contiene códigos UNSPSC válidos para importar.")
        existing = {
            item.code: item
            for item in UNSPSCCode.objects.filter(
                catalog_version=options["catalog_version"], code__in=list(seen)
            )
        }
        created = sum(1 for code in seen if code not in existing)
        updated = sum(
            1 for code, record in seen.items()
            if code in existing and any(
                getattr(existing[code], field) != record[field]
                for field in ("description", "level", "active", "source")
            )
        )
        unchanged = len(seen) - created - updated
        self.stdout.write(
            f"Resumen cambios: nuevos={created}; actualizaciones={updated}; "
            f"sin_cambios={unchanged}; fingerprint={fingerprint}."
        )
        self.stdout.write(f"DRY-RUN UNSPSC: filas_fuente={source_rows}; filas_jerarquia={rows}; códigos_únicos={len(seen)}; repeticiones_jerarquía={duplicates}; inválidos={invalid}; niveles={by_level}; segundos={elapsed}.")
        if not options["apply"]:
            self.stdout.write("No se modificó la base. Usa --apply con confirmación explícita para cargar el catálogo.")
            return
        if options["confirm"] != CONFIRMATION:
            raise CommandError(f'--apply requiere --confirm "{CONFIRMATION}".')
        batch = []
        applied_batches = 0

        def flush(records):
            nonlocal applied_batches
            if not records:
                return
            with transaction.atomic():
                for record in records:
                    UNSPSCCode.objects.update_or_create(
                        catalog_version=record["catalog_version"], code=record["code"],
                        defaults={"description": record["description"], "level": record["level"], "active": record["active"], "source": record["source"]},
                    )
            applied_batches += 1

        for record in seen.values():
            batch.append(record)
            if len(batch) >= batch_size:
                flush(batch)
                batch = []
        flush(batch)
        with transaction.atomic():
            for record in seen.values():
                parent_code = _parent_code(record["code"])
                parent = UNSPSCCode.objects.filter(
                    catalog_version=record["catalog_version"], code=parent_code
                ).first() if parent_code else None
                UNSPSCCode.objects.filter(
                    catalog_version=record["catalog_version"], code=record["code"]
                ).update(parent=parent)
        SyncAuditLog.objects.create(
            operation="import_unspsc_catalog", resource="unspsc", result=SyncAuditLog.RESULT_SUCCESS,
            detail=f"Catálogo UNSPSC {options['catalog_version']}: {len(seen)} códigos actualizados.",
            metadata={"version": options["catalog_version"], "count": len(seen), "duplicates": duplicates, "invalid": invalid, "source": source, "fingerprint": fingerprint, "created": created, "updated": updated, "unchanged": unchanged},
        )
        self.stdout.write(f"Importados/actualizados={len(seen)}; lotes={applied_batches}.")
