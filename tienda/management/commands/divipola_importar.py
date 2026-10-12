"""Carga controlada del catálogo DIVIPOLA suministrado por el administrador."""

import csv
import unicodedata
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from tienda.models import ColombiaMunicipio


CONFIRM = "APLICAR DIVIPOLA"


class Command(BaseCommand):
    help = "Valida o carga un archivo DIVIPOLA CSV/XLSX; dry-run por defecto."

    def add_arguments(self, parser):
        parser.add_argument("archivo")
        parser.add_argument("--catalog-version", dest="catalog_version", required=True)
        parser.add_argument("--apply", action="store_true")
        parser.add_argument("--confirm", default="")

    def handle(self, *args, **options):
        if options["apply"] and options["confirm"] != CONFIRM:
            raise CommandError(f"--apply requiere --confirm \"{CONFIRM}\".")
        rows = list(self._rows(Path(options["archivo"])))
        if not rows:
            raise CommandError("El archivo no contiene municipios válidos.")
        version = options["catalog_version"].strip()
        seen = set()
        valid = []
        invalid = 0
        for row in rows:
            code = str(row.get("codigo") or "").strip().zfill(5)
            dept_code = str(row.get("codigo_departamento") or "").strip().zfill(2)
            name = str(row.get("nombre") or "").strip()
            department = str(row.get("departamento") or "").strip()
            if len(code) != 5 or not code.isdigit() or len(dept_code) != 2 or not dept_code.isdigit() or not name or not department:
                invalid += 1
                continue
            key = (version, code)
            if key in seen:
                invalid += 1
                continue
            seen.add(key)
            valid.append({"codigo": code, "codigo_departamento": dept_code, "nombre": name, "departamento": department})
        if not valid:
            raise CommandError("No hay filas DIVIPOLA válidas después de validar columnas y códigos.")
        codes = {code for _, code in seen}
        existing = {
            row.codigo: row
            for row in ColombiaMunicipio.objects.filter(version=version, codigo__in=codes)
        }
        new_count = sum(code not in existing for code in codes)
        changed = sum(1 for row in valid if row["codigo"] in existing and any(getattr(existing[row["codigo"]], key) != value for key, value in row.items()))
        if options["apply"]:
            with transaction.atomic():
                for row in valid:
                    ColombiaMunicipio.objects.update_or_create(version=version, codigo=row["codigo"], defaults=row)
            self.stdout.write(f"DIVIPOLA aplicada: nuevos={new_count}; actualizados={changed}; invalidos={invalid}.")
        else:
            self.stdout.write(f"DRY-RUN DIVIPOLA: filas={len(rows)}; validos={len(valid)}; nuevos={new_count}; actualizaciones={changed}; invalidos={invalid}.")

    @staticmethod
    def _rows(path):
        if not path.is_file():
            raise CommandError("No se encontró el archivo indicado.")
        suffix = path.suffix.lower()
        if suffix == ".csv":
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle)
                for row in reader:
                    yield Command._normalize_row(row)
            return
        if suffix in {".xlsx", ".xlsm"}:
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise CommandError("El soporte XLSX requiere openpyxl instalado.") from exc
            workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                header_map, first_data_row = Command._find_headers(rows)
                if not header_map:
                    continue
                for values_row in first_data_row:
                    yield Command._normalize_row({key: values_row[index] if index < len(values_row) else "" for key, index in header_map.items()})
            return
        raise CommandError("Formato no admitido. Usa CSV o XLSX.")

    @staticmethod
    def _normalize_row(row):
        aliases = {
            "codigo": "codigo", "codigo municipio": "codigo", "mpio cdpmp": "codigo",
            "municipio": "nombre", "nombre municipio": "nombre", "nombre": "nombre",
            "codigo departamento": "codigo_departamento", "dpto ccdgo": "codigo_departamento",
            "departamento": "departamento", "nombre departamento": "departamento",
        }
        result = {}
        for key, value in row.items():
            normalized = aliases.get(Command._normalize_header(key))
            if normalized:
                result[normalized] = value
        return result

    @staticmethod
    def _normalize_header(value):
        text = unicodedata.normalize("NFKD", str(value or ""))
        text = "".join(char for char in text if not unicodedata.combining(char))
        return " ".join(text.replace("\n", " ").strip().casefold().split())

    @classmethod
    def _find_headers(cls, rows):
        """Find one-row or DANE's two-row hierarchical header."""
        buffered = []
        for row in rows:
            buffered.append(tuple(row or ()))
            if len(buffered) > 40:
                buffered.pop(0)
            if len(buffered) >= 2:
                parent = [cls._normalize_header(value) for value in buffered[-2]]
                child = [cls._normalize_header(value) for value in buffered[-1]]
                if parent[:3] == ["departamento", "", "municipio"] and child[:4] == ["codigo", "nombre", "codigo", "nombre"]:
                    return (
                        {"codigo_departamento": 0, "departamento": 1, "codigo": 2, "nombre": 3},
                        rows,
                    )
            candidate = [cls._normalize_header(value) for value in buffered[-1]]
            mapping = {}
            for index, header in enumerate(candidate):
                canonical = cls._normalize_row({header: ""}).keys()
                for value in canonical:
                    mapping[value] = index
            if {"codigo", "nombre", "codigo_departamento", "departamento"}.issubset(mapping):
                return mapping, rows
        return {}, rows
